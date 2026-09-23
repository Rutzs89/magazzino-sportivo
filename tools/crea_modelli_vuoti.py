"""Crea i quattro fogli Excel **vuoti**, da dare a una societa' nuova.

Sono gli stessi che il programma sa scaricare da Impostazioni, ma averli gia'
pronti in una cartella e' comodo: si copia la cartella, si mandano i fogli ai
dirigenti, e quando tornano indietro si caricano.

Le colonne e le istruzioni devono restare uguali a quelle che il programma si
aspetta: stanno in `MODELLI_EXCEL`, dentro `src/magazzino-sportivo.html`.

Si lancia: python tools/crea_modelli_vuoti.py [cartella]
"""

import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

RADICE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
GIALLO = PatternFill("solid", fgColor="F2C230")

MODELLI = [
    (
        "Articoli",
        ["Articolo", "Ha un numero", "Taglia", "Quantità"],
        [
            "ARTICOLI E MAGAZZINO — da caricare per primo.",
            "",
            "Una riga per ogni articolo e taglia.",
            "",
            "Articolo      il nome, per esempio: TEE NERA ALLENAMENTO, BORSONE U14",
            "Ha un numero  scrivi si solo per le divise da gara, che hanno il numero stampato.",
            "              Per tutto il resto lascia vuoto o scrivi no.",
            "Taglia        XS, S, M, L... oppure UNICA per borsoni, sacche e borracce.",
            "Quantità      quanti pezzi ci sono adesso in magazzino. Le divise con il numero",
            "              non si contano qui: quelle si caricano col foglio Divise.",
        ],
    ),
    (
        "Atleti",
        ["Squadra", "Cognome", "Nome", "Note"],
        [
            "SQUADRE E ATLETE — da caricare dopo gli articoli.",
            "",
            "Una riga per atleta. Le squadre che non esistono vengono create da sole,",
            "nell'ordine in cui compaiono qui: quell'ordine conta anche per le assegnazioni.",
            "",
            "Un'atleta gia' presente nella stessa squadra viene saltata: ricaricare lo",
            "stesso foglio due volte non crea doppioni.",
            "",
            "NON mettere qui mail, telefoni o date delle visite mediche: il programma non",
            "li usa e sono dati di ragazzi e ragazze spesso minorenni.",
        ],
    ),
    (
        "Divise",
        ["Modello", "Taglia", "Numero", "Squadra", "Cognome", "Nome", "Da cambiare", "Note"],
        [
            "DIVISE DA GARA — da caricare dopo gli atleti.",
            "",
            "Una riga per ogni capo fisico.",
            "",
            "Modello       STANDARD per le divise normali. LIBERO per le maglie da libero,",
            "              che valgono per tutte le squadre. Altri nomi creano altri lotti.",
            "Squadra       lascia vuote queste tre se la divisa e' in magazzino. Se ce l'ha",
            "Cognome       un atleta, scrivi la sua squadra, cognome e nome: deve essere",
            "Nome          gia' caricata col foglio Atleti.",
            "Da cambiare   scrivi si se la sta restituendo: il numero torna libero.",
            "",
            "Due maglie con lo stesso modello, taglia e numero in mano ad atlete di squadre",
            "diverse sono due capi distinti, ed e' normale: il numero non si ripete dentro",
            "la stessa squadra, fra squadre si'.",
        ],
    ),
    (
        "Richieste",
        ["Squadra", "Cognome", "Nome", "Articolo", "Taglia", "Numero desiderato", "Note"],
        [
            "RICHIESTE — da caricare per ultimo.",
            "",
            "Una riga per ogni cosa che un atleta deve ricevere. E' il foglio che i",
            "dirigenti di squadra possono compilare e rimandare indietro.",
            "",
            "Numero desiderato   solo per le divise, e solo se ne chiede uno preciso.",
            "",
            "Le righe che non corrispondono a nessun atleta gia' inserita vengono",
            "segnalate e non caricate.",
        ],
    ),
]


def main() -> None:
    dest = sys.argv[1] if len(sys.argv) > 1 else os.path.join(RADICE, "data", "white", "modelli")
    os.makedirs(dest, exist_ok=True)
    for nome, colonne, istruzioni in MODELLI:
        wb = Workbook()
        ws = wb.active
        ws.title = nome
        ws.append(colonne)
        for c, testo in enumerate(colonne, start=1):
            cella = ws.cell(row=1, column=c)
            cella.font = Font(bold=True)
            cella.fill = GIALLO
            ws.column_dimensions[get_column_letter(c)].width = max(12, len(testo) + 4)
        ws.freeze_panes = "A2"

        guida = wb.create_sheet("Istruzioni")
        guida.column_dimensions["A"].width = 100
        for testo in istruzioni:
            guida.append([testo])

        wb.save(os.path.join(dest, f"{nome.lower()}.xlsx"))
        print(f"  {nome.lower()}.xlsx")
    print(f"\nQuattro fogli vuoti in {dest}")


if __name__ == "__main__":
    main()
